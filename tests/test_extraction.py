import io
import zipfile
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from procurement_screening.extraction import extract_document
from procurement_screening.validation import (
    MAX_UPLOAD_BYTES,
    UploadError,
    detect_kind,
    safe_upload_path,
    sanitize_filename,
)


def make_pdf(pages: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        doc.new_page().insert_text((72, 72), text)
    return doc.tobytes()


def make_docx() -> bytes:
    d = Document()
    d.add_paragraph("Bidder must have turnover of Rs 50 lakh.")
    d.add_paragraph("")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "GST"
    t.rows[0].cells[1].text = "Required"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_pages_referenced():
    doc = extract_document("tender.pdf", make_pdf(["First page text", "Second page text"]))
    assert [(b.page, b.text) for b in doc.blocks] == [
        (1, "First page text"),
        (2, "Second page text"),
    ]
    assert doc.blocks[1].locator == "page 2"


def test_pdf_blank_page_skipped():
    doc = extract_document("t.pdf", make_pdf(["", "Only here"]))
    assert [b.page for b in doc.blocks] == [2]


def test_docx_paragraphs_and_tables():
    doc = extract_document("bid.docx", make_docx())
    assert doc.blocks[0].text.startswith("Bidder must have turnover of Rs 50 lakh")
    assert doc.blocks[0].locator == "paragraph 1"
    assert doc.blocks[1].locator == "table 1 row 1"
    assert doc.blocks[1].text == "GST | Required"
    assert all(b.page is None for b in doc.blocks)


@pytest.mark.parametrize(
    "name,data",
    [
        ("x.exe", b"MZ..."),
        ("x.pdf", b"not a pdf"),
        ("x.docx", b"not a docx"),
        ("x.pdf", b""),
        ("x.docx", b"PK\x03\x04garbage"),
    ],
)
def test_rejects_bad_uploads(name, data):
    with pytest.raises(UploadError):
        detect_kind(name, data)


def test_rejects_oversize():
    with pytest.raises(UploadError):
        detect_kind("big.pdf", b"%PDF-" + b"0" * MAX_UPLOAD_BYTES)


def test_rejects_zip_without_document_xml():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(UploadError):
        detect_kind("x.docx", buf.getvalue())


def test_extension_must_match_content():
    with pytest.raises(UploadError):
        detect_kind("sneaky.pdf", make_docx())


def test_corrupt_pdf_rejected():
    with pytest.raises(UploadError):
        extract_document("x.pdf", b"%PDF-1.4 truncated garbage")


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("../../etc/passwd.pdf", "passwd.pdf"),
        ("..\\..\\win\\evil.PDF", "evil.pdf"),
        ("my tender (final).docx", "my_tender_final_.docx"),
        (".hidden.pdf", "hidden.pdf"),
    ],
)
def test_sanitize_filename(raw, expected):
    assert sanitize_filename(raw) == expected


@pytest.mark.parametrize("bad", ["", "noext", "a\x00.pdf", ".pdf"])
def test_sanitize_rejects(bad):
    with pytest.raises(UploadError):
        sanitize_filename(bad)


def test_safe_upload_path_stays_inside(tmp_path: Path):
    p = safe_upload_path(tmp_path, "../../escape.pdf")
    assert p.parent == tmp_path.resolve()
