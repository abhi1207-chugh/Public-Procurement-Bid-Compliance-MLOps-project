"""Extract text from PDF/DOCX with source references.

Every block carries a locator so findings can cite where text came from:
PDFs by 1-based page number; DOCX (which has no fixed pages) by paragraph/table index.
"""
import io

import pymupdf
from docx import Document
from pydantic import BaseModel

from .validation import UploadError, detect_kind

MAX_PDF_PAGES = 300


class TextBlock(BaseModel):
    text: str
    page: int | None = None  # 1-based; None for DOCX
    locator: str  # e.g. "page 3" or "paragraph 12" / "table 1 row 2"


class ExtractedDocument(BaseModel):
    filename: str
    kind: str
    blocks: list[TextBlock]

    @property
    def is_empty(self) -> bool:
        return not self.blocks


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    kind = detect_kind(filename, data)
    blocks = _extract_pdf(data) if kind == "pdf" else _extract_docx(data)
    return ExtractedDocument(filename=filename, kind=kind, blocks=blocks)


def _extract_pdf(data: bytes) -> list[TextBlock]:
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:
        raise UploadError("could not open PDF") from e
    with doc:
        if doc.needs_pass:
            raise UploadError("encrypted PDFs are not supported")
        if doc.page_count > MAX_PDF_PAGES:
            raise UploadError(f"PDF has too many pages (limit {MAX_PDF_PAGES})")
        blocks = []
        for i, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                blocks.append(TextBlock(text=text, page=i, locator=f"page {i}"))
        return blocks


def _extract_docx(data: bytes) -> list[TextBlock]:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise UploadError("could not open DOCX") from e
    blocks = []
    for i, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if text:
            blocks.append(TextBlock(text=text, locator=f"paragraph {i}"))
    for t, table in enumerate(doc.tables, start=1):
        for r, row in enumerate(table.rows, start=1):
            cells = [c.text.strip() for c in row.cells]
            text = " | ".join(c for c in cells if c)
            if text:
                blocks.append(TextBlock(text=text, locator=f"table {t} row {r}"))
    return blocks
